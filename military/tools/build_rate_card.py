#!/usr/bin/env python3
"""
Step 2 of the rate-card pipeline.

Reads data/rate_card.csv (the standardized rate card) and writes:
  * "Wilkins Military Rate Card - MASTER.xlsx"  - the formatted Excel master David's team maintains
  * data/rate_card.js                            - the same data embedded for the proposal builder app

Re-run after editing rate_card.csv.  (The app can also load the MASTER .xlsx directly, so the
Excel file can be edited in place and re-uploaded without running this script.)
"""
import csv, json, os, datetime, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image as XLImage

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, '..')
CSV = os.path.join(ROOT, 'data', 'rate_card.csv')
XLSX = os.path.join(ROOT, 'Wilkins Military Rate Card - MASTER.xlsx')
JS = os.path.join(ROOT, 'data', 'rate_card.js')
LOGO = os.path.join(ROOT, 'assets', 'wilkins-logo.png')

NAVY = '1F2A44'; BLUE = '0092E4'; LIGHT = 'EDF2F6'; WHITE = 'FFFFFF'; GREY = '6B7280'
THIN = Side(style='thin', color='D1D5DB')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FONT = Font(name='Calibri', bold=True, color=WHITE, size=11)
HDR_FILL = PatternFill('solid', fgColor=NAVY)
SUB_FILL = PatternFill('solid', fgColor=BLUE)
BAND_FILL = PatternFill('solid', fgColor='F8FAFC')
MONEY = '"$"#,##0;[Red]-"$"#,##0;"-"'

SETTINGS = [
    ('Standard margin (share of client price kept by Wilkins)', 0.50, 'Client price = Vendor cost / (1 - margin).  50% = vendor x 2, the standard military markup.  Competitive bids have used 30%.'),
    ('Round client prices to nearest', 25, 'Matches the rounding used in the prior master list ($25).'),
    ('Standard production - Poster 22"W x 28"H (client price)', 100, 'Per David, Sep 22 2026 (raised because frames must ship separately).'),
    ('Standard production - Banner 3\'H x 6\'W (client price)', 350, 'Per David, Sep 22 2026.'),
    ('Est. printer cost - Poster', 50, 'Estimate; update from the printer quote.'),
    ('Est. printer cost - Banner 3x6', 96, 'Printer quote (Cecily): $95.51 with ground shipping.'),
    ('Default campaign length (months)', 3, 'A one-month campaign is not as effective; pricing is quoted on a three-month campaign.'),
    ('Standard pricing method', 'Monthly rate x months', 'House rule: use the 1-month rate x months (not the discounted term price) unless it is a competitive bid.'),
    ('Rate card year considered current', 2025, 'Cards from 2025/2026 are good; older cards should be reconfirmed with the base.'),
]

COLUMNS = [
    # (csv key, header, width, number format)
    ('id', 'ID', 18, None), ('base', 'Base', 30, None), ('aka', 'Also known as', 22, None), ('state', 'State', 8, None),
    ('dma', 'DMA / Metro', 24, None), ('zip', 'Zip', 8, None), ('branch', 'Branch', 14, None),
    ('category', 'Category', 18, None), ('media_type', 'Media Type', 28, None), ('ad_unit_size', 'Ad Unit Size', 18, None),
    ('description', 'Description', 50, None), ('min_term_months', 'Min Term (mo)', 10, None),
    ('vendor_1mo', 'Vendor Cost 1 Month', 14, MONEY), ('vendor_3mo', 'Vendor Cost 3 Months', 14, MONEY),
    ('vendor_6mo', 'Vendor Cost 6 Months', 14, MONEY), ('vendor_12mo', 'Vendor Cost 12 Months', 14, MONEY),
    ('vendor_production', 'Vendor Production Cost', 14, MONEY),
    ('client_1mo', 'Client Price 1 Month', 14, MONEY), ('client_3mo', 'Client Price 3 Months', 14, MONEY),
    ('client_6mo', 'Client Price 6 Months', 14, MONEY), ('client_12mo', 'Client Price 12 Months', 14, MONEY),
    ('client_production', 'Client Production Price', 14, MONEY),
    ('active_duty', 'Active Duty (ref.)', 12, '#,##0'), ('dependents', 'Dependents (ref.)', 12, '#,##0'),
    ('rate_card_year', 'Rate Card Year', 10, None), ('last_updated', 'Last Updated', 12, None),
    ('notes', 'Notes', 40, None), ('review', 'Review Flag', 30, None), ('source', 'Source', 30, None),
]

def num(v):
    try:
        return float(v) if v not in ('', None) else None
    except ValueError:
        return None

def style_header(ws, row, ncols, fill=HDR_FILL, height=42):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR_FONT; cell.fill = fill; cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[row].height = height

def build_xlsx(rows):
    wb = Workbook()
    # ---------- Read Me ----------
    ws = wb.active; ws.title = 'Read Me'
    ws.sheet_view.showGridLines = False
    if os.path.exists(LOGO):
        img = XLImage(LOGO); img.width, img.height = 218, 48; ws.add_image(img, 'B2')
    ws['B6'] = 'Wilkins Media - Military On-Base Advertising Rate Card (MASTER)'; ws['B6'].font = Font(bold=True, size=16, color=NAVY)
    ws['B7'] = 'Generated %s from the standardized rate card (data/rate_card.csv).' % datetime.date.today().strftime('%B %d, %Y'); ws['B7'].font = Font(color=GREY, italic=True)
    lines = [
        ('What this is', 'One row per media option on each installation, with the VENDOR (base) cost at 1 / 3 / 6 / 12 months and production. Client prices are formulas driven by the Settings tab, so changing the margin re-prices every line.'),
        ('Source', 'Box > Military > Master Grid > "MasterMilitaryOnBaseList 12.2.25 NS.xlsx", cleaned and normalized. Base populations are reference values from "Military - National Digital Inventory Master 6.3.26.xlsx".'),
        ('How to update a price', 'Edit the Vendor Cost cells on the Rate Card tab. Put the year of the rate card you used in "Rate Card Year" and today\'s date in "Last Updated". Do not type into the Client Price columns - they are formulas.'),
        ('How to add a base or option', 'Add a row at the bottom of the Rate Card table (the table extends automatically). Fill Base, State, DMA, Branch, Category, Media Type, Size, Description and the vendor costs you know. Leave unknown terms blank.'),
        ('Pricing conventions', 'Vendor Cost columns are TOTALS for that term (3 Months = total for a 3-month flight). Where a base only sells a longer term, the 1 Month column is blank and Min Term shows the shortest term sold. Marquees and other per-unit items are priced per unit unless the description says otherwise.'),
        ('Client price rule', 'Client Price = Vendor Cost / (1 - margin), rounded to the nearest $25 (Settings tab). Standard military margin is 50% (vendor x 2). Competitive bids may use a lower margin - change it in the proposal builder, not here.'),
        ('Production', 'Production is never included in the media price. Standard production (quantity of 1): Poster 22x28 $100, Banner 3x6 $350. Other sizes: inquire. Where a base quoted its own production, it is in the Vendor Production Cost column.'),
        ('Review Flag', '"No price on file" = the base lists the item but we have no cost (inquire). "Rate card older than 2025" = reconfirm with the base before quoting. "Source formula errors" = the old master had #VALUE! in that row; verify against the base rate card in Box > Military > 2 Military Installations.'),
        ('Proposal builder', 'The web app (military/index.html in the EMC_AI_Tool repo) reads this file: use "Load rate card" in the app to pull in your latest edits, then build the proposal and export the client and internal Excel files.'),
    ]
    r = 9
    for k, v in lines:
        ws.cell(row=r, column=2, value=k).font = Font(bold=True, color=NAVY)
        c = ws.cell(row=r, column=3, value=v); c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 48; r += 1
    ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 26; ws.column_dimensions['C'].width = 110

    # ---------- Settings ----------
    st = wb.create_sheet('Settings'); st.sheet_view.showGridLines = False
    st['A1'] = 'Setting'; st['B1'] = 'Value'; st['C1'] = 'Notes'; style_header(st, 1, 3, height=22)
    for i, (k, v, n) in enumerate(SETTINGS, start=2):
        st.cell(row=i, column=1, value=k).border = BORDER
        c = st.cell(row=i, column=2, value=v); c.border = BORDER; c.font = Font(bold=True, color=BLUE)
        if isinstance(v, float): c.number_format = '0%'
        elif isinstance(v, int) and i in (3, 4, 5, 6, 7): c.number_format = MONEY
        st.cell(row=i, column=3, value=n).border = BORDER
        st.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical='top')
    st.column_dimensions['A'].width = 52; st.column_dimensions['B'].width = 24; st.column_dimensions['C'].width = 90

    # ---------- Rate Card ----------
    rc = wb.create_sheet('Rate Card')
    for ci, (k, h, w, fmt) in enumerate(COLUMNS, start=1):
        rc.cell(row=1, column=ci, value=h); rc.column_dimensions[get_column_letter(ci)].width = w
    style_header(rc, 1, len(COLUMNS))
    col = {k: i for i, (k, h, w, f) in enumerate(COLUMNS, start=1)}
    L = lambda k: get_column_letter(col[k])
    for ri, row in enumerate(rows, start=2):
        for k, h, w, fmt in COLUMNS:
            ci = col[k]
            if k.startswith('client_'):
                src = 'vendor_' + k.split('_', 1)[1]
                v = '=IF(ISNUMBER(%s%d),ROUND(%s%d/(1-Settings!$B$2)/Settings!$B$3,0)*Settings!$B$3,"")' % (L(src), ri, L(src), ri)
            elif k in ('vendor_1mo', 'vendor_3mo', 'vendor_6mo', 'vendor_12mo', 'vendor_production', 'active_duty', 'dependents', 'min_term_months'):
                v = num(row[k])
                if v is not None and v == int(v): v = int(v)
            else:
                v = row[k]
            c = rc.cell(row=ri, column=ci, value=v); c.border = BORDER
            c.alignment = Alignment(wrap_text=k in ('description', 'notes', 'review', 'media_type', 'base'), vertical='top')
            if fmt: c.number_format = fmt
            if k.startswith('client_'): c.fill = PatternFill('solid', fgColor='EEF7FD'); c.font = Font(color=NAVY, bold=True)
            if k == 'review' and row[k]: c.font = Font(color='B45309')
    tab = Table(displayName='RateCard', ref='A1:%s%d' % (get_column_letter(len(COLUMNS)), len(rows) + 1))
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
    rc.add_table(tab)
    rc.freeze_panes = 'C2'

    # ---------- Bases ----------
    bs = wb.create_sheet('Bases')
    heads = ['Base', 'Also known as', 'State', 'DMA / Metro', 'Zip', 'Branch', 'Media options', 'Digital', 'Standard Media', 'Large Format', 'Print & Pubs', 'Web & Email', 'Active Duty (ref.)', 'Dependents (ref.)', 'Rate card year(s)', 'Rows needing review']
    for ci, h in enumerate(heads, start=1): bs.cell(row=1, column=ci, value=h)
    style_header(bs, 1, len(heads))
    bases = collections.OrderedDict()
    for r in rows:
        b = bases.setdefault(r['base'], dict(aka=r['aka'], state=r['state'], dma=r['dma'], zip=r['zip'], branch=r['branch'], n=0, cats=collections.Counter(), ad=r['active_duty'], dep=r['dependents'], years=set(), review=0))
        b['n'] += 1; b['cats'][r['category']] += 1
        if r['rate_card_year']: b['years'].add(r['rate_card_year'])
        if r['review']: b['review'] += 1
    for ri, (name, b) in enumerate(sorted(bases.items(), key=lambda kv: (kv[1]['state'], kv[0])), start=2):
        vals = [name, b['aka'], b['state'], b['dma'], b['zip'], b['branch'], b['n'], b['cats']['Digital'], b['cats']['Standard Media'], b['cats']['Large Format'], b['cats']['Print & Publications'], b['cats']['Web & Email'], num(b['ad']), num(b['dep']), ', '.join(sorted(b['years'])), b['review']]
        for ci, v in enumerate(vals, start=1):
            c = bs.cell(row=ri, column=ci, value=v); c.border = BORDER
            if ci in (13, 14): c.number_format = '#,##0'
    for ci, w in enumerate([32, 26, 7, 26, 8, 14, 10, 8, 10, 10, 10, 10, 12, 12, 14, 12], start=1):
        bs.column_dimensions[get_column_letter(ci)].width = w
    bs.freeze_panes = 'B2'
    bs.add_table(Table(displayName='Bases', ref='A1:%s%d' % (get_column_letter(len(heads)), len(bases) + 1), tableStyleInfo=TableStyleInfo(name='TableStyleLight9', showRowStripes=True)))

    # ---------- Change Log ----------
    cl = wb.create_sheet('Change Log')
    for ci, h in enumerate(['Date', 'Who', 'Base', 'What changed', 'Rate card / source used'], start=1): cl.cell(row=1, column=ci, value=h)
    style_header(cl, 1, 5, height=22)
    cl.cell(row=2, column=1, value=datetime.date.today().isoformat()); cl.cell(row=2, column=2, value='Claude (for Jon)')
    cl.cell(row=2, column=3, value='All'); cl.cell(row=2, column=4, value='Initial standardized master built from MasterMilitaryOnBaseList 12.2.25 NS.xlsx; client prices converted to margin-driven formulas; categories assigned; review flags added.')
    cl.cell(row=2, column=5, value='Box > Military > Master Grid')
    for ci, w in enumerate([12, 18, 30, 80, 40], start=1): cl.column_dimensions[get_column_letter(ci)].width = w
    wb.save(XLSX)
    print('wrote', XLSX)

def build_js(rows):
    keys = ['id','base','aka','state','dma','zip','branch','category','media_type','ad_unit_size','description',
            'vendor_1mo','vendor_3mo','vendor_6mo','vendor_12mo','vendor_production','min_term_months',
            'active_duty','dependents','rate_card_year','last_updated','notes','review']
    data = []
    for r in rows:
        o = {}
        for k in keys:
            v = r[k]
            if k.startswith('vendor_') or k in ('active_duty', 'dependents', 'min_term_months'):
                v = num(v)
                if v is not None and v == int(v): v = int(v)
            if v in ('', None): continue
            o[k] = v
        data.append(o)
    settings = {
        'margin': 0.50, 'roundTo': 25, 'prodPosterClient': 100, 'prodBannerClient': 350,
        'prodPosterCost': 50, 'prodBannerCost': 96, 'defaultMonths': 3, 'currentYear': 2025,
    }
    payload = {'generated': datetime.date.today().isoformat(), 'source': 'Wilkins Military Rate Card - MASTER.xlsx (from MasterMilitaryOnBaseList 12.2.25 NS.xlsx)', 'settings': settings, 'rows': data}
    with open(JS, 'w') as f:
        f.write('// Generated by tools/build_rate_card.py - do not edit by hand. Edit the MASTER .xlsx / rate_card.csv and rebuild.\n')
        f.write('window.WILKINS_RATE_CARD = ' + json.dumps(payload, separators=(',', ':')) + ';\n')
    print('wrote', JS, len(data), 'rows')

if __name__ == '__main__':
    rows = list(csv.DictReader(open(CSV)))
    build_xlsx(rows)
    build_js(rows)
