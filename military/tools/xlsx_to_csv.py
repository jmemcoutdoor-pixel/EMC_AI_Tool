#!/usr/bin/env python3
"""
Turn a MASTER workbook (the "Rate Card" tab, as downloaded from the app or edited in Excel)
back into data/rate_card.csv so the rest of the pipeline can rebuild the app from it.

    python3 tools/xlsx_to_csv.py "Wilkins Military Rate Card - MASTER.xlsx"

Client Price columns are formulas in the workbook and are ignored; vendor cost is the source of truth.
"""
import csv, os, re, sys, warnings
warnings.filterwarnings('ignore')
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.abspath(os.path.join(HERE, '..'))
OUT = os.path.join(ROOT, 'data', 'rate_card.csv')
HEAD = {  # csv column -> header regex on the Rate Card tab
    'id': r'^id$', 'base': r'^base$', 'aka': r'also known', 'state': r'^state', 'dma': r'dma', 'zip': r'^zip', 'branch': r'^branch', 'category': r'^category',
    'media_type': r'media type', 'ad_unit_size': r'ad unit size', 'description': r'^description', 'min_term_months': r'min term',
    'vendor_1mo': r'vendor cost 1 month', 'vendor_3mo': r'vendor cost 3 month', 'vendor_6mo': r'vendor cost 6 month', 'vendor_12mo': r'vendor cost 12 month', 'vendor_production': r'vendor production',
    'active_duty': r'active duty', 'dependents': r'dependents', 'rate_card_year': r'rate card year', 'last_updated': r'last updated', 'notes': r'^notes', 'review': r'review', 'source': r'^source$',
    'box_folder_url': r'box folder', 'box_source_file': r'rate card file in box|box source file', 'box_source_url': r'box source url|link to rate card',
}
COLS = list(HEAD.keys())

def main(path):
    wb = load_workbook(path, data_only=False)
    ws = wb['Rate Card'] if 'Rate Card' in wb.sheetnames else wb[wb.sheetnames[0]]
    hdr = [str(c.value or '').strip() for c in ws[1]]
    idx = {}
    for k, rx in HEAD.items():
        for i, h in enumerate(hdr):
            if re.search(rx, h, re.I): idx[k] = i; break
    if 'base' not in idx or 'media_type' not in idx: sys.exit('No Base / Media Type columns found on the Rate Card tab')
    out = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        vals = [c.value for c in row]
        get = lambda k: vals[idx[k]] if k in idx and idx[k] < len(vals) else ''
        base, media = str(get('base') or '').strip(), str(get('media_type') or '').strip()
        if not base or not media: continue
        r = {}
        for k in COLS:
            v = get(k)
            if k.startswith('vendor_') or k in ('active_duty', 'dependents', 'min_term_months'):
                try: v = float(v) if v not in (None, '') else ''
                except (TypeError, ValueError): v = ''
                if v != '' and v == int(v): v = int(v)
            elif k in ('box_folder_url', 'box_source_url') and k in idx and idx[k] < len(row) and row[idx[k]].hyperlink:
                v = row[idx[k]].hyperlink.target or v
            r[k] = '' if v is None else str(v).strip() if not isinstance(v, (int, float)) else v
        if r['min_term_months'] == '':
            for t, k in ((1, 'vendor_1mo'), (3, 'vendor_3mo'), (6, 'vendor_6mo'), (12, 'vendor_12mo')):
                if r[k] not in ('', 0): r['min_term_months'] = t; break
        out.append(r)
    with open(OUT, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=COLS); w.writeheader(); w.writerows(out)
    print('wrote', OUT, len(out), 'rows from', os.path.basename(path))

if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'Wilkins Military Rate Card - MASTER.xlsx'))
